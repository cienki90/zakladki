"""Odczyt pliku Excel 'zestawienie' do listy wpisow (entries).

Kolumny lokalizowane sa po NAZWIE naglowka (bez rozroznienia wielkosci liter),
a nie po stalym indeksie - w arkuszu wzorcowym kolumna D jest pusta,
a 'wt' znajduje sie w kolumnie E.
"""

from dataclasses import dataclass
from typing import List

from openpyxl import load_workbook

# Nazwy wymaganych naglowkow (porownanie bez wielkosci liter).
REQUIRED_HEADERS = ("trafo", "miejscowosc", "wt")
PREFERRED_SHEET = "Arkusz1"


@dataclass
class Entry:
    """Pojedynczy wpis zakladki."""

    trafo: str
    miejscowosc: str
    wt: str


def _norm(value) -> str:
    """Zamienia wartosc komorki naglowka na znormalizowany tekst."""
    if value is None:
        return ""
    return str(value).strip().lower()


def _cell_to_str(value) -> str:
    """Zamienia wartosc komorki danych na tekst, zachowujac dokladne wartosci.

    openpyxl zwraca liczby jako int/float - zachowujemy je jako tekst, ale
    unikamy '.0' dla liczb calkowitych zapisanych jako float.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_entries(path) -> List[Entry]:
    """Czyta plik Excel i zwraca liste wpisow {trafo, miejscowosc, wt}.

    Args:
        path: sciezka do pliku .xlsx.

    Returns:
        Lista obiektow Entry dla kazdego niepustego wiersza danych.

    Raises:
        ValueError: gdy brakuje ktoregokolwiek z wymaganych naglowkow.
    """
    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    try:
        if PREFERRED_SHEET in wb.sheetnames:
            ws = wb[PREFERRED_SHEET]
        else:
            ws = wb[wb.sheetnames[0]]

        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            header = ()

        # Mapowanie: znormalizowana nazwa naglowka -> indeks kolumny.
        header_index = {}
        for idx, cell in enumerate(header):
            name = _norm(cell)
            if name and name not in header_index:
                header_index[name] = idx

        missing = [h for h in REQUIRED_HEADERS if h not in header_index]
        if missing:
            raise ValueError(
                "Brak wymaganych kolumn w arkuszu: "
                + ", ".join(missing)
                + ". Wymagane naglowki (wiersz 1): trafo, miejscowosc, wt."
            )

        i_trafo = header_index["trafo"]
        i_miej = header_index["miejscowosc"]
        i_wt = header_index["wt"]

        entries: List[Entry] = []
        for row in rows:
            if row is None:
                continue
            trafo = _cell_to_str(row[i_trafo]) if i_trafo < len(row) else ""
            miejscowosc = _cell_to_str(row[i_miej]) if i_miej < len(row) else ""
            wt = _cell_to_str(row[i_wt]) if i_wt < len(row) else ""

            # Pomijamy calkowicie puste wiersze.
            if not trafo and not miejscowosc and not wt:
                continue

            entries.append(Entry(trafo=trafo, miejscowosc=miejscowosc, wt=wt))

        return entries
    finally:
        wb.close()
